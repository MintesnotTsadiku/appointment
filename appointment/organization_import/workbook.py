"""Versioned organization workbook contract and bounded, write-free parsing."""

import hashlib
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError, fromstring
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from frappe import _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

VERSION = "appointment-organization-workbook.v1"
MAX_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 30 * 1024 * 1024
MAX_ROWS = 2000
SHEETS = {
    "Organization": ("key", "name", "timezone", "email", "phone", "description"),
    "Locations": ("key", "name", "timezone", "address", "phone"),
    "Team": ("key", "email", "display_name", "role", "provider_key", "location_keys"),
    "Providers": ("key", "name", "email", "description"),
    "Services": ("key", "name", "duration", "price", "location_key", "provider_key", "capacity", "public"),
    "Availability": ("key", "location_key", "weekday", "opens_at", "closes_at", "timezone"),
    "Website Content": ("key", "field", "text"),
}
REQUIRED = {
    "Organization": {"key", "name", "timezone"}, "Locations": {"key", "name", "timezone"},
    "Team": {"key", "email", "display_name", "role"}, "Providers": {"key", "name", "email"},
    "Services": {"key", "name", "duration", "location_key", "provider_key", "capacity", "public"},
    "Availability": {"key", "location_key", "weekday", "opens_at", "closes_at", "timezone"},
    "Website Content": {"key", "field", "text"},
}
DAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")


def template():
    book = Workbook()
    instructions = book.active
    instructions.title = "Read Me"
    for line in (VERSION, "Fill the named sheets; Examples is never imported.",
                 "Keep each stable key unchanged when correcting or retrying an import.",
                 "Use existing staff email addresses. Never enter passwords or formulas.",
                 "Use IANA time zones, Monday–Sunday, and 24-hour HH:MM times.",
                 "Upload for a dry run, review each cell error, then explicitly confirm.",
                 "Team invitations require separate confirmation; this workbook does not send email."):
        instructions.append([line])
    instructions.column_dimensions["A"].width = 100
    for name, fields in SHEETS.items():
        sheet = book.create_sheet(name)
        sheet.append(fields)
        sheet.freeze_panes = "A2"
        for index, column in enumerate(fields, 1):
            sheet.cell(1, index).font = Font(bold=True)
            sheet.column_dimensions[get_column_letter(index)].width = 24
    examples = book.create_sheet("Examples")
    examples.append(["Examples only — these rows are never imported"])
    examples.append(["Organization", "clinic", "Example Clinic", "Africa/Addis_Ababa"])
    examples.append(["Locations", "main", "Main office", "Africa/Addis_Ababa"])
    examples.append(["Team", "owner", "owner@example.test", "Example Owner", "Manager"])
    examples.append(["Availability", "main-mon", "main", "Monday", "09:00", "17:00", "Africa/Addis_Ababa"])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def dry_run(content):
    """Return normalized proposed rows and cell errors; never access a database."""
    errors = []
    data = {name: [] for name in SHEETS}
    try:
        _check_archive(content)
        book = load_workbook(BytesIO(content), read_only=False, data_only=False, keep_links=False)
    except (ValueError, BadZipFile, KeyError, OSError, ParseError, InvalidFileException):
        return {"version": VERSION, "valid": False, "errors": [_error("Read Me", 1, "A", _("Choose a valid current XLSX template without macros, external links, unsafe XML, or oversized archive entries."))], "rows": data}
    try:
        if "Read Me" not in book.sheetnames or book["Read Me"]["A1"].value != VERSION:
            errors.append(_error("Read Me", 1, "A", _("Download the current workbook template.")))
        unknown = set(book.sheetnames) - set(SHEETS) - {"Read Me", "Examples"}
        for name in sorted(unknown):
            errors.append(_error(name, 1, "A", _("Remove this unsupported sheet.")))
        for name, fields in SHEETS.items():
            if name not in book.sheetnames:
                errors.append(_error(name, 1, "A", _("Restore this sheet from the current template.")))
                continue
            _read_sheet(book[name], fields, data[name], errors)
        _references(data, errors)
        if len(data["Organization"]) != 1:
            errors.append(_error("Organization", 2, "A", _("Provide exactly one organization row.")))
        return {"version": VERSION, "valid": not errors, "errors": errors, "rows": data,
                "sha256": hashlib.sha256(content).hexdigest(),
                "summary": {name: len(rows) for name, rows in data.items()}}
    finally:
        book.close()


def _check_archive(content):
    if not isinstance(content, bytes) or not content or len(content) > MAX_BYTES:
        raise ValueError("Choose an XLSX workbook smaller than 5 MB.")
    with ZipFile(BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > 200 or len({item.filename for item in members}) != len(members):
            raise ValueError("This workbook archive contains too many or duplicate entries.")
        if sum(item.file_size for item in members) > MAX_EXPANDED_BYTES:
            raise ValueError("This workbook expands beyond the 30 MB limit.")
        for item in members:
            path = item.filename.lower()
            if item.flag_bits & 1 or path.startswith("/") or ".." in path.split("/") or "\\" in path:
                raise ValueError("This workbook contains an unsafe archive entry.")
            if any(part in path for part in ("vbaproject", "externallinks", "embeddings", "activex")):
                raise ValueError("Remove macros, embedded files, and external links from the workbook.")
            if item.file_size > max(1024 * 1024, item.compress_size * 200):
                raise ValueError("This workbook contains an excessively compressed entry.")
            if path.endswith((".xml", ".rels")):
                # Restrict to the current template's UTF-8 XML before parsing;
                # alternate encodings cannot hide declarations from this scan.
                raw = archive.read(item).decode("utf-8-sig")
                xml = raw.upper()
                if "\x00" in raw or "<!DOCTYPE" in xml or "<!ENTITY" in xml:
                    raise ValueError("This workbook contains unsupported XML declarations.")
                if 'TARGETMODE="EXTERNAL"' in xml or "TARGETMODE='EXTERNAL'" in xml:
                    raise ValueError("Remove external relationships from the workbook.")
                root = fromstring(raw)
                if any(key.lower().endswith("targetmode") and value.lower() == "external"
                       for node in root.iter() for key, value in node.attrib.items()):
                    raise ValueError("Remove external relationships from the workbook.")


def _error(sheet, row, column, message):
    return {"sheet": sheet, "row": row, "cell": f"{column}{row}", "message": message}


def _read_sheet(sheet, fields, rows, errors):
    if sheet.merged_cells.ranges:
        errors.append(_error(sheet.title, 1, "A", _("Unmerge cells and restore the template layout.")))
        return
    if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > len(fields):
        errors.append(_error(sheet.title, 1, "A", _("Use the template columns and at most 2,000 rows per sheet.")))
        return
    headers = tuple(cell.value for cell in sheet[1])
    if headers != fields:
        errors.append(_error(sheet.title, 1, "A", _("Keep the template column headings in their original order.")))
        return
    keys = set()
    for cells in sheet.iter_rows(min_row=2):
        if all(cell.value is None for cell in cells):
            continue
        row = {"_row": cells[0].row}
        for field, cell in zip(fields, cells):
            value = str(cell.value).strip() if cell.value is not None else ""
            row[field] = value
            if cell.data_type == "f" or value.startswith(("=", "@")) or (value.startswith("+") and field != "phone") or cell.hyperlink:
                errors.append(_error(sheet.title, cell.row, cell.column_letter, _("Replace formulas and links with plain values.")))
            elif len(value) > 2400 or any(ord(character) < 32 for character in value if character not in "\n\r\t"):
                errors.append(_error(sheet.title, cell.row, cell.column_letter, _("Use plain text of at most 2,400 characters.")))
            elif not value and field in REQUIRED[sheet.title]:
                errors.append(_error(sheet.title, cell.row, cell.column_letter, _("Fill in {0}.").format(field.replace("_", " "))))
            elif value:
                message = _validate(field, value, sheet.title)
                if message:
                    errors.append(_error(sheet.title, cell.row, cell.column_letter, message))
        if row["key"] in keys:
            errors.append(_error(sheet.title, cells[0].row, "A", _("Use a different stable key for each row.")))
        keys.add(row["key"])
        rows.append(row)


def _validate(field, value, sheet):
    if any(character in value for character in "<>"):
        return _("Use plain text without HTML markup.")
    if field == "name" and len(value) > 100:
        return _("Keep names to at most 100 characters.")
    if field == "description" and len(value) > 500:
        return _("Keep descriptions to at most 500 characters.")
    if field.endswith("key") and not re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", value):
        return _("Use a stable key with 1–64 letters, digits, underscores, or hyphens.")
    if field == "email" and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value):
        return _("Enter a complete email address.")
    if field == "timezone":
        try:
            ZoneInfo(value)
        except (ZoneInfoNotFoundError, ValueError):
            return _("Enter an IANA time zone, such as Africa/Addis_Ababa.")
    if field in {"duration", "capacity"}:
        ceiling = 480 if field == "duration" else 100
        floor = 5 if field == "duration" else 1
        if not value.isdigit() or not floor <= int(value) <= ceiling:
            return _("Enter a whole number between {0} and {1}.").format(floor, ceiling)
    if field == "price":
        if not re.fullmatch(r"\d{1,7}(\.\d{1,2})?", value):
            return _("Enter a non-negative price with at most two decimal places.")
    if field == "public" and value not in {"0", "1"}:
        return _("Enter 1 for public or 0 for private.")
    if field == "role" and value not in {"Manager", "Provider", "Receptionist"}:
        return _("Choose Manager, Provider, or Receptionist.")
    if field == "weekday" and value not in DAYS:
        return _("Use a full weekday name, such as Monday.")
    if field in {"opens_at", "closes_at"} and not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", value):
        return _("Enter a 24-hour time in HH:MM format.")
    if sheet == "Website Content" and field == "field" and value not in {"hero_title", "hero_subtitle", "about_body", "contact_email", "contact_phone"}:
        return _("Choose hero_title, hero_subtitle, about_body, contact_email, or contact_phone.")
    return None


def _references(data, errors):
    keys = {name: {row["key"] for row in rows} for name, rows in data.items()}
    days = set()
    text_fields = set()
    for sheet, rows in data.items():
        for row in rows:
            for field, target in (("location_key", "Locations"), ("provider_key", "Providers")):
                if row.get(field) and row[field] not in keys[target]:
                    column = get_column_letter(SHEETS[sheet].index(field) + 1)
                    errors.append(_error(sheet, row["_row"], column, _("Use a key from the {0} sheet.").format(target)))
            if sheet == "Team":
                for key in filter(None, (part.strip() for part in row["location_keys"].split(","))):
                    if key not in keys["Locations"]:
                        errors.append(_error(sheet, row["_row"], "F", _("Use comma-separated keys from Locations.")))
            if sheet == "Availability" and row["opens_at"] >= row["closes_at"]:
                errors.append(_error(sheet, row["_row"], "E", _("Closing time must follow opening time on the same day.")))
            if sheet == "Availability":
                identity = (row["location_key"], row["weekday"])
                if identity in days:
                    errors.append(_error(sheet, row["_row"], "C", _("Provide one hours row per location and weekday.")))
                days.add(identity)
            if sheet == "Website Content":
                field = row["field"]
                if field in text_fields:
                    errors.append(_error(sheet, row["_row"], "B", _("Provide one row per website text field.")))
                text_fields.add(field)
                limit = {"hero_title": 1000, "hero_subtitle": 1000, "about_body": 2400, "contact_email": 160, "contact_phone": 32}.get(field, 2400)
                if len(row["text"]) > limit:
                    errors.append(_error(sheet, row["_row"], "C", _("Keep this text to at most {0} characters.").format(limit)))
