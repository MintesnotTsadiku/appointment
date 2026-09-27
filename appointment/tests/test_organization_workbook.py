"""Malformed workbook and cell-level dry-run regression tests; no database."""

from io import BytesIO
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from appointment.organization_import import workbook


class WorkbookTests(unittest.TestCase):
    def content(self, change=None):
        book = load_workbook(BytesIO(workbook.template()))
        book["Organization"].append(["clinic", "Example clinic", "Africa/Addis_Ababa", "owner@example.test", "+251911123456", ""])
        book["Locations"].append(["main", "Main office", "Africa/Addis_Ababa", "Example address", ""])
        book["Providers"].append(["owner", "Example owner", "owner@example.test", ""])
        book["Team"].append(["owner", "owner@example.test", "Example owner", "Provider", "owner", "main"])
        book["Services"].append(["visit", "Consultation", 30, 0, "main", "owner", 1, 0])
        book["Availability"].append(["main-mon", "main", "Monday", "09:00", "17:00", "Africa/Addis_Ababa"])
        if change:
            change(book)
        stream = BytesIO()
        book.save(stream)
        book.close()
        return stream.getvalue()

    def test_valid_workbook_and_examples_are_not_imported(self):
        result = workbook.dry_run(self.content())
        self.assertTrue(result["valid"], result["errors"])
        self.assertEqual(result["summary"]["Organization"], 1)
        self.assertNotIn("Examples", result["rows"])
        self.assertEqual(result["rows"]["Organization"][0]["phone"], "+251911123456")

    def test_formula_is_a_plain_cell_error(self):
        result = workbook.dry_run(self.content(lambda book: setattr(book["Services"]["C2"], "value", "=10+20")))
        self.assertFalse(result["valid"])
        self.assertTrue(any(row["cell"] == "C2" and "formulas" in row["message"] for row in result["errors"]))

    def test_missing_reference_is_located(self):
        result = workbook.dry_run(self.content(lambda book: setattr(book["Services"]["E2"], "value", "foreign")))
        self.assertTrue(any(row["sheet"] == "Services" and row["cell"] == "E2" for row in result["errors"]))

    def test_duplicate_stable_key_is_rejected(self):
        result = workbook.dry_run(self.content(lambda book: book["Locations"].append(["main", "Second office", "Africa/Addis_Ababa"])))
        self.assertTrue(any(row["cell"] == "A3" and "stable key" in row["message"] for row in result["errors"]))

    def test_wrong_version_and_password_column_are_rejected(self):
        def change(book):
            book["Read Me"]["A1"] = "old-template"
            book["Team"]["G1"] = "password"
            book["Team"]["G2"] = "never import"
        result = workbook.dry_run(self.content(change))
        self.assertFalse(result["valid"])
        self.assertTrue(any(row["sheet"] == "Read Me" for row in result["errors"]))
        self.assertTrue(any(row["sheet"] == "Team" for row in result["errors"]))

    def test_invalid_timezone_and_reversed_hours_are_reported(self):
        def change(book):
            book["Locations"]["C2"] = "Mars/Office"
            book["Availability"]["E2"] = "08:00"
        result = workbook.dry_run(self.content(change))
        self.assertTrue(any("IANA" in row["message"] for row in result["errors"]))
        self.assertTrue(any("Closing" in row["message"] for row in result["errors"]))

    def test_oversized_row_range_and_merged_cells_are_rejected(self):
        def change(book):
            book["Locations"]["A2002"] = "too-many"
            book["Team"].merge_cells("A2:B2")
        result = workbook.dry_run(self.content(change))
        self.assertTrue(any("2,000" in row["message"] for row in result["errors"]))
        self.assertTrue(any("Unmerge" in row["message"] for row in result["errors"]))

    def test_non_archive_and_external_links_are_rejected(self):
        self.assertFalse(workbook.dry_run(b"not a workbook")["valid"])
        content = BytesIO(self.content())
        with ZipFile(content, "a") as archive:
            archive.writestr("xl/externalLinks/externalLink1.xml", "<link/>")
        self.assertFalse(workbook.dry_run(content.getvalue())["valid"])

    def test_entity_archive_is_rejected_before_xml_parsing(self):
        content = BytesIO()
        with ZipFile(content, "w") as archive:
            archive.writestr("xl/workbook.xml", '<!DOCTYPE workbook [<!ENTITY x "unsafe">]><workbook/>')
        self.assertFalse(workbook.dry_run(content.getvalue())["valid"])
        encoded = BytesIO()
        with ZipFile(encoded, "w") as archive:
            archive.writestr("xl/workbook.xml", '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE workbook [<!ENTITY x "unsafe">]><workbook>&x;</workbook>'.encode("utf-16"))
        self.assertFalse(workbook.dry_run(encoded.getvalue())["valid"])

    def test_broken_xml_and_oversized_input_return_plain_errors(self):
        content = BytesIO()
        with ZipFile(content, "w") as archive:
            archive.writestr("xl/workbook.xml", "<workbook>")
        self.assertFalse(workbook.dry_run(content.getvalue())["valid"])
        self.assertFalse(workbook.dry_run(b"x" * (workbook.MAX_BYTES + 1))["valid"])


if __name__ == "__main__":
    unittest.main()
